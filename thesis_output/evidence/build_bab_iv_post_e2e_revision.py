from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = REPO_ROOT / "thesis_output" / "REVISI_BAB_IV_PASCA_E2E_20260825.docx"
PHONE_HOME = REPO_ROOT / "artifacts" / "sidang_live" / "phone_home.png"
PHONE_E2E = REPO_ROOT / "artifacts" / "sidang_live" / "phone_e2e_visual.png"
BACKEND_IMAGE = REPO_ROOT / "artifacts" / "sidang_live" / "e2e_backend_padi.png"

# Required document design preset: narrative_proposal. Typography and page geometry
# are deliberately overridden to match the user's existing thesis document.
DESIGN_PRESET = "narrative_proposal"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=12, bold=None, italic=None, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=12, bold=None, italic=None, color=None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)


def add_body(document: Document, text: str, *, first_line=True, italic=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(1.25) if first_line else Cm(0)
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run(text)
    set_run_font(run, size=12, italic=italic)
    return paragraph


def add_heading(document: Document, text: str, *, level=1):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True)
    return paragraph


def add_list_item(document: Document, number: int, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.75)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{number}.\t{text}")
    set_run_font(run, size=12)
    return paragraph


def add_caption(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=11, bold=True)
    return paragraph


def style_table(table, widths_cm: list[float] | None = None, font_size=9.5) -> None:
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths_cm and column_index < len(widths_cm):
                cell.width = Cm(widths_cm[column_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_paragraph_font(
                    paragraph,
                    size=font_size,
                    bold=True if row_index == 0 else None,
                    color="FFFFFF" if row_index == 0 else None,
                )
            if row_index == 0:
                set_cell_shading(cell, "2E6F52")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(table, widths_cm)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=10)


def add_source_note(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=10, italic=True)


def build_document() -> None:
    for required in (PHONE_HOME, PHONE_E2E, BACKEND_IMAGE):
        if not required.is_file():
            raise FileNotFoundError(required)

    document = Document()
    document.core_properties.title = "Revisi BAB IV Pasca Pengujian End-to-End TaniAdaptif"
    document.core_properties.subject = "Naskah pengganti BAB IV berdasarkan bukti repository per 25 Agustus 2026"
    document.core_properties.author = "TaniAdaptif"
    document.core_properties.keywords = "TaniAdaptif, BAB IV, Stable Diffusion, LoRA, Flutter, E2E"

    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.top_margin = Cm(4)
    section.bottom_margin = Cm(3)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Cm(3)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("REVISI BAB IV PASCA PENGUJIAN END-TO-END")
    set_run_font(run, size=14, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    run = subtitle.add_run("TaniAdaptif")
    set_run_font(run, size=14, bold=True, color="2E6F52")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Cm(2)
    run = meta.add_run("Naskah pengganti dan bagian sinkronisasi\nStatus bukti: 25 Agustus 2026")
    set_run_font(run, size=12)

    note_table = document.add_table(rows=1, cols=1)
    note_table.autofit = False
    note_cell = note_table.cell(0, 0)
    note_cell.width = Cm(14)
    set_cell_shading(note_cell, "EAF4EE")
    set_cell_margins(note_cell, top=150, start=180, bottom=150, end=180)
    note_cell.text = (
        "Dokumen ini tidak mengubah file skripsi utama. Bagian 4.3 di bawah dapat "
        "menggantikan draf lama secara utuh. Bagian 4.1.8, 4.1.9, dan 4.2 disediakan "
        "sebagai sinkronisasi karena status E2E dan riwayat telah berubah. Nomor tabel "
        "dan gambar perlu disesuaikan dengan naskah utama."
    )
    for paragraph in note_cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.15
        set_paragraph_font(paragraph, size=11)

    document.add_page_break()

    add_heading(document, "A. Ringkasan Perubahan Utama")
    add_body(
        document,
        "Revisi ini menghapus batas klaim pra-E2E yang sudah tidak sesuai dengan kondisi aktual. "
        "Jalur Flutter Android ke Express, FastAPI, GPU, Stable Diffusion 1.5, LoRA final, PNG, "
        "dan kembali ke perangkat telah dijalankan. Bukti perangkat tersimpan di repository. "
        "Namun, keberhasilan tersebut tetap dibatasi sebagai bukti fungsional pada lingkungan "
        "pengujian, bukan bukti efektivitas agronomi atau kesiapan produksi.",
    )

    add_table(
        document,
        ["Aspek", "Draf lama", "Revisi berdasarkan bukti aktual"],
        [
            [
                "Inference E2E",
                "Belum diuji",
                "Lulus pada lingkungan pengujian: Android -> Express -> FastAPI Diffusers/CUDA + LoRA -> PNG -> Android.",
            ],
            [
                "Bukti Android",
                "Belum tersedia",
                "Tersedia phone_e2e_visual.png (720 x 1600) dan phone_home.png (720 x 1600).",
            ],
            [
                "Riwayat",
                "Implementasi parsial; layar belum tersedia",
                "RiwayatScreen dan navigasi tersedia; penyimpanan dibatasi 20 entri; widget test lulus.",
            ],
            [
                "Regresi perangkat lunak",
                "96 ML, 29 inference, 23 backend, 13 Flutter pada 21 Agustus",
                "Dijalankan ulang 25 Agustus: 96 ML pipeline, 37 inference inti, 26 backend, dan 18 Flutter; seluruhnya lulus.",
            ],
            [
                "Variasi gambar",
                "Seed tetap menyebabkan hasil identik atau sangat mirip",
                "Kontrak lokal telah diubah: seed eksplisit tetap reproducible; seed kosong menjadi acak; filename unik. Sinkronisasi ke runtime Kaggle aktif belum diklaim.",
            ],
        ],
        [3.2, 4.3, 6.5],
    )

    add_heading(document, "B. Naskah Pengganti Bagian 4.3")
    add_body(
        document,
        "Salin bagian berikut untuk menggantikan 4.3 sampai 4.3.8 pada draf lama. Sitasi [3], "
        "[4], [5], [6], [12], [13], [14], [17], [19], dan [23] dipertahankan agar dapat "
        "diselaraskan dengan daftar pustaka naskah utama.",
        first_line=False,
        italic=True,
    )

    add_heading(document, "4.3 Pembahasan")
    add_heading(document, "4.3.1 Ketercapaian Tujuan Penelitian", level=2)
    add_body(
        document,
        "Tujuan pertama adalah merancang dan membangun aplikasi mobile yang mengintegrasikan "
        "data cuaca BMKG, kondisi lahan, dan rule-based engine. Tujuan ini tercapai pada tingkat "
        "implementasi serta pengujian perangkat lunak dan perangkat Android. Aplikasi Flutter "
        "dapat menampilkan cuaca wilayah Aikmel melalui backend Express, menerima input kondisi "
        "lahan, menampilkan rekomendasi beserta rule ID, menyediakan kalender empat minggu, dan "
        "menyimpan hasil ke riwayat. Backend menormalisasi respons BMKG, mempertahankan nilai "
        "curah hujan yang tidak tersedia sebagai null, serta memproses 30 kunci aturan secara "
        "deterministik.",
    )
    add_body(
        document,
        "Tujuan kedua adalah menerapkan fine-tuning LoRA pada Stable Diffusion 1.5 untuk "
        "menghasilkan simulasi visual empat komoditas. Smoke test 100 langkah dan main training "
        "1.600 langkah selesai pada GPU cloud. Model final dari run 20260820T143551Z bernama "
        "multisource_multicrop_lora.safetensors, berukuran 6.414.992 byte, memuat 256 tensor, "
        "dan mempunyai SHA-256 0bda8908d577513110c22cf5ec1b8a7e008d2b26dd90525732a5c4e4b1bdd43a. "
        "Model tersebut selanjutnya dimuat pada runtime GPU dan digunakan melalui FastAPI untuk "
        "menghasilkan PNG yang diteruskan Express hingga tampil pada perangkat Android. Dengan "
        "demikian, tujuan kedua tidak lagi hanya terbukti pada tahap training, tetapi juga pada "
        "jalur inference end-to-end di lingkungan pengujian.",
    )
    add_body(
        document,
        "Tujuan ketiga adalah menguji fungsionalitas aplikasi dan mengevaluasi model secara "
        "teknis. Pengujian regresi yang dijalankan ulang pada 25 Agustus 2026 menghasilkan 96 "
        "test ML pipeline, 37 test inference inti, 26 test backend, dan 18 test Flutter yang "
        "seluruhnya lulus. Evaluasi manusia menggunakan tepat 72 penilaian atau 36 pasangan "
        "Base-LoRA, dengan sembilan kondisi pada setiap komoditas serta prompt dan seed identik "
        "dalam setiap pasangan. Oleh karena itu, tujuan ketiga tercapai dalam ruang lingkup "
        "pengujian teknis penelitian. Ketercapaian ini tidak diperluas menjadi klaim efektivitas "
        "agronomi lapangan, penerimaan pengguna, atau kesiapan produksi.",
    )

    add_heading(document, "4.3.2 Pembahasan Integrasi BMKG", level=2)
    add_body(
        document,
        "Integrasi cuaca menjawab kebutuhan agar data BMKG tidak hanya ditampilkan, tetapi juga "
        "diproses sebagai input sistem. Pendekatan ini mempunyai keterkaitan dengan penelitian "
        "aplikasi cuaca berbasis Android untuk pertanian [3] dan pemanfaatan data BMKG [14]. "
        "Perbedaan implementasi TaniAdaptif terletak pada normalisasi data, cache dengan status "
        "freshness, fallback saat jaringan gagal, serta penerusan kategori hujan ke rule-based "
        "engine. Uji live wilayah Aikmel pada 5 Agustus 2026 memproses 18 titik prakiraan untuk "
        "tiga hari dan seluruh 18 nilai curah hujan tersedia. Screenshot perangkat pada 24 "
        "Agustus 2026 juga memperlihatkan cuaca Aikmel, sumber BMKG, suhu, kelembapan, angin, "
        "dan kategori hujan pada halaman beranda.",
    )
    add_body(
        document,
        "Horizon prakiraan BMKG yang singkat tetap menjadi batas metodologis. Sistem tidak "
        "mempunyai dasar untuk menyatakan prakiraan bulanan resmi. Proyeksi empat minggu dibentuk "
        "sebagai kalender simulasi berdasarkan ringkasan kondisi yang tersedia, sedangkan minggu "
        "kedua sampai keempat mempunyai tingkat keyakinan lebih rendah. Dengan demikian, kalender "
        "harus dipahami sebagai rencana rotasi berbasis kondisi saat permintaan, bukan prediksi "
        "iklim jangka panjang.",
    )

    add_heading(document, "4.3.3 Pembahasan Rule-Based Recommendation", level=2)
    add_body(
        document,
        "Rule-based engine memisahkan fungsi keputusan dari model generatif. Pemisahan ini "
        "mendukung keterlacakan karena setiap rekomendasi mempunyai kunci aturan, alasan, dan "
        "rule ID. Pendekatan tersebut berada dalam konteks sistem pendukung keputusan pertanian "
        "yang dibahas pada penelitian kesesuaian lahan [4] dan pemilihan tanaman [12], tetapi "
        "TaniAdaptif menggunakan aturan deterministik yang dirancang untuk empat komoditas, bukan "
        "metode AHP, SAW, atau Fuzzy Weighted Product. Stable Diffusion dan LoRA tidak menentukan "
        "tanaman; keduanya hanya memvisualisasikan tanaman yang telah dipilih oleh aturan.",
    )
    add_body(
        document,
        "Pengujian 30 kunci membuktikan kelengkapan kontrak perangkat lunak, bukan kebenaran "
        "agronomi universal. Aturan hanya menggunakan kategori cuaca, ketersediaan air, tanaman "
        "sebelumnya, dan jenis lahan. Sistem tidak menggunakan hasil laboratorium tanah, varietas, "
        "hama, harga pasar, atau prediksi hasil panen. Karena itu, rekomendasi merupakan keluaran "
        "prototipe simulasi rotasi tanam dalam batas penelitian.",
    )
    add_body(
        document,
        "Masih terdapat perbedaan antara Tabel 3.6 pada naskah utama dan implementasi final. "
        "Tabel 3.6 memuat 12 contoh dengan istilah basah dan kering, sedangkan kode final memakai "
        "30 kunci dengan kategori tinggi, sedang, dan rendah. Perbedaan tersebut harus "
        "diselaraskan pada BAB III agar uraian metode dan hasil merujuk pada decision table yang "
        "sama.",
    )

    add_heading(document, "4.3.4 Pembahasan Implementasi Aplikasi", level=2)
    add_body(
        document,
        "Implementasi Flutter memperlihatkan bahwa alur cuaca, input kondisi lahan, rekomendasi, "
        "kalender, simulasi visual, dan riwayat dapat ditempatkan dalam satu aplikasi Android. "
        "Hal ini sejalan secara tematik dengan penggunaan Android pada sistem pertanian [13] dan "
        "landasan pengembangan mobile dengan Flutter [17]. TaniAdaptif menambahkan cache berbasis "
        "signature, isolasi kegagalan visual, dan proxy Express agar token inference privat serta "
        "alamat upstream tidak disimpan di aplikasi.",
    )
    add_body(
        document,
        "Pengujian menunjukkan bahwa kegagalan jaringan atau layanan visual tidak menghilangkan "
        "rekomendasi yang telah dihasilkan rule-based engine. Layar Riwayat Rekomendasi juga "
        "telah tersedia dan dapat memuat data lokal, sedangkan penyimpanan riwayat dibatasi pada "
        "20 entri terbaru. Pada pengujian perangkat nyata, gambar hasil inference berhasil tampil "
        "pada layar Simulasi Visual. Meskipun demikian, penelitian ini tidak melakukan UAT, "
        "wawancara, kuesioner, maupun pengujian dengan responden petani. Oleh sebab itu, hasil "
        "implementasi tidak digunakan untuk menyatakan kemudahan penggunaan atau efektivitas pada "
        "pengguna akhir.",
    )

    add_heading(document, "4.3.5 Pembahasan Fine-Tuning dan Evaluasi LoRA", level=2)
    add_body(
        document,
        "Fine-tuning LoRA menerapkan konsep penyesuaian model generatif yang juga dibahas pada "
        "penelitian Stable Diffusion untuk karakter visual khusus [5], penerapan LoRA pada Stable "
        "Diffusion [6], dan landasan latent diffusion [19]. Perbandingan terhadap penelitian "
        "tersebut bersifat konseptual karena dataset, domain, prompt, dan metriknya berbeda. Hasil "
        "training dan loss finite membuktikan pipeline optimisasi mencapai artefak final, tetapi "
        "loss yang fluktuatif tidak digunakan sebagai bukti langsung kualitas visual.",
    )
    add_body(
        document,
        "Penilaian manusia pada 36 pasangan menunjukkan kesesuaian komoditas sebesar 31 dari 36 "
        "atau 86,11% untuk Base dan 34 dari 36 atau 94,44% untuk LoRA. Selisih deskriptifnya "
        "adalah 8,33 poin persentase. Exact McNemar menghasilkan p=0,25 sehingga selisih tersebut "
        "tidak signifikan pada taraf 0,05. Median kualitas konteks dan realisme keseluruhan turun "
        "dari 4,5 pada Base menjadi 4,0 pada LoRA. Exact sign test menghasilkan p=0,6072387695. "
        "Hasil ini menunjukkan bahwa LoRA tidak meningkatkan seluruh dimensi visual secara "
        "seragam.",
    )
    add_body(
        document,
        "Peningkatan identitas terutama terlihat pada kacang tanah: kesesuaian kelas berubah dari "
        "44,44% pada Base menjadi 77,78% pada LoRA, sedangkan median konteks dan realisme berubah "
        "dari 2 menjadi 4. Sebaliknya, jagung tetap dikenali 100% pada kedua varian tetapi median "
        "konteks dan realisme turun dari 5 pada Base menjadi 4 pada LoRA. Karena setiap kelas "
        "hanya mempunyai sembilan pasangan dan uji per kelas bersifat eksploratif, hasil tersebut "
        "tidak cukup untuk generalisasi di luar prompt, seed, dataset, serta empat kelas penelitian.",
    )

    add_heading(document, "4.3.6 Pembahasan Pengujian dan Keandalan Bukti", level=2)
    add_body(
        document,
        "Black Box Testing digunakan sesuai landasan pada BAB II [23] dengan memeriksa kontrak "
        "input-output, status, dan penanganan error. Suite otomatis mencakup parser dan cache BMKG, "
        "30 aturan rekomendasi, kalender, validasi inference, retry, timeout, proxy, penyimpanan "
        "lokal, serta widget Flutter. Validator artefak menambah pemeriksaan CRC ZIP, hash model, "
        "struktur SafeTensors, manifest, dan keterbacaan PNG. Pengujian ulang pada 25 Agustus "
        "2026 menghasilkan 96 test ML pipeline, 37 test inference inti, 26 test backend, dan 18 "
        "test Flutter yang seluruhnya lulus tanpa menjalankan Stable Diffusion pada laptop.",
    )
    add_body(
        document,
        "Pemisahan antara mock dan inference nyata tetap diperlukan. Provider mock hanya "
        "membuktikan kontrak API pada CPU, sedangkan uji end-to-end nyata membuktikan rangkaian "
        "Flutter Android, adb reverse ke Express, HTTPS Cloudflare, FastAPI pada GPU, Stable "
        "Diffusion 1.5 dengan LoRA final, pengiriman PNG melalui Express, dan penampilan gambar di "
        "perangkat. Screenshot Android membuktikan jalur fungsional dan rendering antarmuka, tetapi "
        "tidak membuktikan ketepatan semantik setiap gambar. Penilaian semantik tetap merujuk pada "
        "36 pasangan evaluasi yang ditinjau manusia.",
    )
    add_body(
        document,
        "Uji end-to-end juga menemukan bahwa penggunaan seed tetap 42 pada permintaan aplikasi "
        "menyebabkan hasil untuk komoditas dan konteks yang sama menjadi identik atau sangat mirip. "
        "Perbaikan lokal telah menerapkan seed opsional: seed eksplisit, termasuk 42, tetap dipakai "
        "untuk reproduktibilitas pengujian; jika seed tidak dikirim atau bernilai null, service "
        "membuat seed baru dan mengembalikannya pada metadata. Identitas file juga menggunakan "
        "crop, seed, dan request ID agar artefak tidak bertabrakan. Regression test lokal untuk "
        "perilaku ini telah lulus, tetapi sinkronisasi perubahan ke snapshot runtime Kaggle aktif "
        "dan pemasangan ulang APK belum dijadikan bukti pada dokumen ini.",
    )

    add_heading(document, "4.3.7 Keterbatasan Penelitian", level=2)
    add_body(document, "Keterbatasan penelitian yang masih berlaku adalah sebagai berikut.")
    limitations = [
        "BMKG menyediakan horizon prakiraan jangka pendek; proyeksi empat minggu merupakan simulasi rencana dan bukan prakiraan bulanan resmi.",
        "Rule-based engine dibatasi pada empat komoditas dan empat kelompok variabel; hasilnya bukan prediksi hasil panen atau validasi agronomi lapangan.",
        "Decision table pada BAB III belum sepenuhnya sama dengan 30 kunci pada implementasi final dan masih memerlukan penyelarasan terminologi.",
        "Dataset berjumlah 389 gambar dan sumber lokal Aikmel hanya terdapat pada sebagian kelas padi, sedangkan tiga kelas lain berasal dari sumber publik.",
        "Evaluasi visual dilakukan oleh peneliti pada 36 pasangan; identitas penilai tidak diverifikasi independen dan tidak ada penilai ahli atau responden pengguna yang diklaim.",
        "Setiap kelas hanya mempunyai sembilan pasangan dan beberapa uji per kelas bersifat eksploratif, sehingga hasil tidak mendukung generalisasi luas.",
        "Metrik brightness, contrast, dan entropy hanya merupakan proxy teknis serta tidak mengukur ketepatan komoditas atau ketepatan agronomi.",
        "Host inference menggunakan sesi GPU cloud dan Cloudflare Quick Tunnel yang bersifat sementara. URL dapat berubah dan layanan berhenti ketika sesi berakhir, sehingga arsitektur ini belum merupakan deployment produksi.",
        "Perbaikan variasi seed telah tervalidasi pada source dan test lokal, tetapi belum dibuktikan telah tersinkron pada runtime Kaggle aktif. Karena itu, deployment yang sedang berjalan dapat masih menghasilkan gambar berulang untuk input yang sama.",
        "Screenshot Android yang tersedia membuktikan pengiriman dan penampilan gambar end-to-end, bukan kesesuaian semantik seluruh empat komoditas pada setiap permintaan.",
        "APK masih debug-signed dan belum menjadi artefak rilis produksi. Penelitian juga belum mencakup UAT atau pengujian lapangan dengan petani.",
    ]
    for index, limitation in enumerate(limitations, start=1):
        add_list_item(document, index, limitation)

    add_heading(document, "4.3.8 Ringkasan Pembahasan Pasca-E2E", level=2)
    add_body(
        document,
        "Hasil penelitian membuktikan bahwa aplikasi Android, backend Express, integrasi BMKG, "
        "rule-based engine, pipeline training, model LoRA, evaluasi visual, dan inference GPU telah "
        "diintegrasikan dalam satu jalur teknis end-to-end. Rekomendasi tetap berasal dari aturan "
        "yang dapat dilacak, sedangkan Stable Diffusion 1.5 dan LoRA berfungsi sebagai modul "
        "visualisasi. Model LoRA menunjukkan peningkatan deskriptif kesesuaian komoditas, terutama "
        "pada kacang tanah, tetapi perbedaannya tidak signifikan pada exact McNemar dan tidak "
        "meningkatkan seluruh skor visual secara seragam.",
    )
    add_body(
        document,
        "Keberhasilan E2E mengubah status penelitian dari prototipe dengan kontrak mock menjadi "
        "prototipe yang telah menjalankan inference nyata pada GPU dan menampilkan hasil pada "
        "Android. Meskipun demikian, sifat sesi cloud yang sementara, belum adanya UAT dan "
        "validasi agronomi lapangan, serta sinkronisasi deployment untuk variasi seed tetap menjadi "
        "batas klaim. Dengan batas tersebut, hasil penelitian dapat dinyatakan memenuhi tujuan "
        "teknis yang ditetapkan tanpa menganggap sistem sebagai layanan produksi atau alat "
        "pengambilan keputusan agronomi yang telah tervalidasi di lapangan.",
    )

    document.add_page_break()
    add_heading(document, "C. Sinkronisasi Bagian 4.1 dan 4.2")
    add_body(
        document,
        "Bagian berikut perlu ikut diganti agar tidak bertentangan dengan pembahasan pasca-E2E. "
        "Teks ini bukan bagian baru, melainkan pengganti paragraf lama yang masih memuat label "
        "[BUKTI E2E GPU BELUM TERSEDIA], [SCREENSHOT ANDROID BELUM TERSEDIA], dan pernyataan "
        "bahwa layar riwayat belum ada.",
        first_line=False,
        italic=True,
    )

    add_heading(document, "4.1.8 Hasil Implementasi Simulasi Visual", level=2)
    add_body(
        document,
        "FastAPI menyediakan endpoint health, readiness, model info, generation, validasi input, "
        "timeout, retry terbatas, penguncian satu proses generasi, validasi hash LoRA, dan "
        "pemeriksaan CUDA. Express menyimpan token inference privat, menerapkan pembatasan akses, "
        "menyembunyikan URL upstream, dan melayani PNG melalui domain backend. Pada uji live 24 "
        "Agustus 2026, aplikasi Flutter pada perangkat Android mengirim permintaan melalui adb "
        "reverse ke Express. Express meneruskan permintaan melalui HTTPS Cloudflare ke FastAPI pada "
        "GPU Tesla T4. Runtime memuat Stable Diffusion 1.5 dan LoRA final dengan SHA-256 "
        "0bda8908d577513110c22cf5ec1b8a7e008d2b26dd90525732a5c4e4b1bdd43a, menghasilkan "
        "PNG 512 x 512, lalu mengembalikan gambar melalui Express hingga tampil pada layar "
        "Simulasi Visual Android.",
    )
    add_body(
        document,
        "Hasil tersebut membuktikan integrasi fungsional end-to-end pada lingkungan pengujian. "
        "Bukti disimpan sebagai artifacts/sidang_live/phone_e2e_visual.png dan "
        "artifacts/sidang_live/e2e_backend_padi.png. Bukti ini tidak digunakan untuk menyatakan "
        "bahwa setiap gambar selalu tepat secara semantik. Kualitas visual tetap dinilai melalui "
        "evaluasi Base-LoRA yang terpisah.",
    )

    add_heading(document, "4.1.9 Hasil Penyimpanan Riwayat", level=2)
    add_body(
        document,
        "Penyimpanan lokal menggunakan SharedPreferences untuk cache cuaca, cache rekomendasi, "
        "dan riwayat visual. Setiap entri riwayat memuat request ID, tanaman, URL gambar backend, "
        "seed, model, prompt, waktu, dan ringkasan konteks. Daftar dibatasi pada 20 entri terbaru. "
        "Layar Riwayat Rekomendasi dan navigasinya telah tersedia. Pengujian widget membuktikan "
        "halaman riwayat dapat menampilkan detail yang disimpan, sedangkan pengujian service "
        "membuktikan bahwa penyimpanan 25 fixture mempertahankan 20 entri terbaru. Screenshot "
        "beranda memperlihatkan akses Riwayat pada navigasi bawah, tetapi screenshot khusus daftar "
        "riwayat belum disertakan sebagai bukti terpisah.",
    )

    add_heading(document, "Pembaruan Paragraf 4.2 Pengujian Sistem", level=2)
    add_body(
        document,
        "Tabel 4.12 memisahkan pengujian kontrak dengan provider mock dari pengujian inference "
        "nyata. Baris mock tetap berstatus Lulus (mock) karena hanya membuktikan kontrak API. "
        "Setelah uji perangkat pada 24 Agustus 2026, baris inference E2E GPU dan tampilan Android "
        "dapat diperbarui menjadi lulus dalam lingkungan pengujian. Klaim tersebut didukung oleh "
        "PNG backend dan screenshot perangkat, tetapi tidak dipakai sebagai bukti kualitas "
        "semantik seluruh output atau kesiapan produksi.",
    )

    add_caption(document, "Tabel 4.X Pembaruan Baris Pengujian Pasca-E2E")
    add_table(
        document,
        ["No.", "Fitur", "Hasil aktual", "Status"],
        [
            [
                "21",
                "Inference E2E GPU",
                "Flutter Android -> Express -> FastAPI Diffusers pada CUDA + LoRA final -> PNG 512 x 512 -> Express berhasil dijalankan.",
                "Lulus",
            ],
            [
                "22",
                "Tampilan Android nyata",
                "Gambar hasil LoRA termuat pada layar Simulasi Visual; bukti phone_e2e_visual.png berukuran 720 x 1600 tersedia.",
                "Lulus (fungsional)",
            ],
            [
                "Tambahan",
                "Variasi simulasi",
                "Seed opsional dan filename unik lulus pada regression test lokal; belum dibuktikan tersinkron ke runtime Kaggle aktif.",
                "Lulus lokal; deployment belum divalidasi",
            ],
        ],
        [1.0, 3.0, 7.5, 2.5],
    )
    add_source_note(
        document,
        "Sumber: artefak sidang live, source lokal, dan pengujian regresi TaniAdaptif, 2026.",
    )

    add_caption(document, "Tabel 4.Y Ringkasan Pengujian Regresi Terbaru")
    add_table(
        document,
        ["Lapisan", "Cakupan", "Hasil", "Status"],
        [
            ["ML pipeline", "Dataset, training contract, artefak, human review", "96 passed", "Lulus"],
            ["Inference inti", "API, E2E probe lokal, validasi LoRA", "37 passed", "Lulus"],
            ["Express backend", "BMKG, rule engine, proxy visual, seed contract", "26 passed", "Lulus"],
            ["Flutter", "DTO, service, cache, visual, riwayat, widget", "18 passed", "Lulus"],
        ],
        [3.0, 6.5, 2.2, 2.3],
    )
    add_source_note(document, "Sumber: pengujian ulang lokal pada 25 Agustus 2026.")

    document.add_page_break()
    add_heading(document, "D. Bukti Visual yang Tersedia")
    add_body(
        document,
        "Gambar berikut dapat dipindahkan ke BAB IV dengan nomor yang disesuaikan. Screenshot "
        "menyimpan bukti antarmuka aktual; gambar tidak dimodifikasi selain penyesuaian ukuran di "
        "dokumen ini.",
        first_line=False,
    )

    figure_table = document.add_table(rows=2, cols=2)
    figure_table.autofit = False
    for row in figure_table.rows:
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for cell in figure_table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=0, start=50, bottom=0, end=50)
        cell.width = Cm(6.8)
    p1 = figure_table.cell(0, 0).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run().add_picture(str(PHONE_HOME), width=Cm(5.5))
    p2 = figure_table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run().add_picture(str(PHONE_E2E), width=Cm(5.5))
    figure_table.cell(1, 0).text = "(a) Beranda dengan data cuaca BMKG Aikmel"
    figure_table.cell(1, 1).text = "(b) Hasil inference nyata pada layar Simulasi Visual"
    for cell in figure_table.rows[1].cells:
        set_cell_margins(cell, top=60, start=50, bottom=60, end=50)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            set_paragraph_font(paragraph, size=10, italic=True)
    add_caption(document, "Gambar 4.X Implementasi TaniAdaptif pada Perangkat Android")
    add_source_note(document, "Sumber: dokumentasi pengujian perangkat TaniAdaptif, 24 Agustus 2026.")
    add_body(
        document,
        "Catatan interpretasi: panel (b) membuktikan gambar hasil provider Diffusers dan LoRA "
        "dikirim sampai ke aplikasi. Satu screenshot tidak dipakai untuk menggeneralisasi "
        "ketepatan semantik semua kelas.",
        first_line=False,
        italic=True,
    )

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.add_run().add_picture(str(BACKEND_IMAGE), width=Cm(8.0))
    add_caption(document, "Gambar 4.Y Artefak PNG 512 x 512 dari Uji Backend End-to-End")
    add_source_note(document, "Sumber: artifacts/sidang_live/e2e_backend_padi.png, 2026.")

    add_heading(document, "E. Daftar Bukti Repository")
    add_table(
        document,
        ["Bukti", "Path relatif", "Kegunaan"],
        [
            ["Screenshot beranda", "artifacts/sidang_live/phone_home.png", "Data BMKG dan navigasi Android"],
            ["Screenshot E2E", "artifacts/sidang_live/phone_e2e_visual.png", "Gambar inference tampil di Android"],
            ["PNG backend", "artifacts/sidang_live/e2e_backend_padi.png", "Artefak PNG 512 x 512"],
            ["Analisis review", "artifacts/model_evaluation/kaggle_run_20260820T143551Z/human_review_analysis.json", "72 penilaian dan statistik Base-LoRA"],
            ["Laporan review", "artifacts/model_evaluation/kaggle_run_20260820T143551Z/POST_REVIEW_ANALYSIS_REPORT.md", "Metode, hasil, hash, dan batas klaim"],
            ["Model final", "artifacts/models/multisource_multicrop_lora_sd15/multisource_multicrop_lora.safetensors", "LoRA final terverifikasi"],
        ],
        [3.2, 7.2, 3.6],
    )

    add_heading(document, "F. Catatan Sebelum Ditempel ke Naskah Utama")
    notes = [
        "Hapus semua placeholder [BUKTI E2E GPU BELUM TERSEDIA] dan [SCREENSHOT ANDROID BELUM TERSEDIA] yang sudah digantikan oleh bukti aktual.",
        "Sesuaikan nomor Gambar 4.X, Gambar 4.Y, Tabel 4.X, dan Tabel 4.Y dengan urutan pada naskah utama.",
        "Perbarui Tabel 3.6 agar kategori dan jumlah aturan sama dengan implementasi final sebelum sidang.",
        "Jangan menulis bahwa variasi gambar acak sudah aktif di Kaggle sebelum source runtime disinkronkan, APK dibangun ulang, dan dua permintaan tanpa seed dibuktikan menghasilkan seed berbeda.",
        "Pertahankan seed 42 pada evaluasi teknis berpasangan karena diperlukan untuk reproduktibilitas; penggunaan seed acak hanya untuk permintaan simulasi baru dari aplikasi.",
        "Jangan memakai screenshot E2E sebagai pengganti hasil human review. Screenshot membuktikan jalur sistem, sedangkan tabel evaluasi membahas kualitas visual.",
    ]
    for index, note in enumerate(notes, start=1):
        add_list_item(document, index, note)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
