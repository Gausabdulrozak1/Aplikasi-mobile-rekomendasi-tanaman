from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "docs" / "BAB_IV_FINAL_DRAFT.md"
DEFAULT_OUTPUT = ROOT / "thesis_output" / "BAB_IV_DRAFT_PRE_E2E.docx"

THESIS_TITLE = (
    "RANCANG BANGUN APLIKASI MOBILE BERBASIS STABLE DIFFUSION DENGAN "
    "FINE-TUNING LoRA DAN API BMKG UNTUK SIMULASI ROTASI TANAM ADAPTIF "
    "IKLIM DI WILAYAH AIKMEL"
)

INLINE_MARKUP = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
TABLE_CAPTION = re.compile(r"^\*\*(Tabel 4\.(\d+) .+?)\*\*$")
FIGURE_CAPTION = re.compile(r"^\*(Gambar 4\.(\d+) .+?)\*$")
IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)$")
ORDERED_ITEM = re.compile(r"^(\d+)\.\s+(.+)$")

WIDE_TABLES = {6, 7, 9, 10, 12}

TABLE_WIDTHS_CM: dict[int, list[float]] = {
    1: [3.0, 6.5, 4.5],
    2: [4.5, 9.5],
    3: [3.4, 4.6, 6.0],
    4: [4.4, 3.2, 3.2, 3.2],
    5: [5.0, 9.0],
    6: [3.2, 2.3, 2.4, 2.4, 3.9, 3.2, 3.8],
    7: [2.1, 1.8, 2.1, 2.1, 2.1, 2.1, 2.1, 2.1, 2.6, 2.6],
    8: [5.0, 3.0, 3.0, 3.0],
    9: [3.0, 2.0, 2.4, 2.4, 4.1, 4.1, 3.0],
    10: [4.0, 3.7, 2.5, 2.0, 2.5, 2.2, 3.4],
    11: [3.0, 2.0, 3.0, 3.0, 3.0],
    12: [0.8, 2.0, 3.2, 2.7, 4.4, 5.6, 1.6],
    13: [3.3, 4.4, 4.3, 2.0],
}


@dataclass(frozen=True)
class Block:
    kind: str
    value: object
    meta: object | None = None


def parse_markdown(source: Path) -> list[Block]:
    lines = source.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    index = 0

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        table_caption = TABLE_CAPTION.match(line)
        if table_caption:
            blocks.append(
                Block("table_caption", table_caption.group(1), int(table_caption.group(2)))
            )
            index += 1
            continue

        figure_caption = FIGURE_CAPTION.match(line)
        if figure_caption:
            blocks.append(
                Block("figure_caption", figure_caption.group(1), int(figure_caption.group(2)))
            )
            index += 1
            continue

        image = IMAGE.match(line)
        if image:
            blocks.append(Block("image", image.group(2), image.group(1)))
            index += 1
            continue

        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [split_table_row(item) for item in table_lines]
            if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
                raise ValueError(f"Tabel Markdown tanpa separator valid di sekitar baris {index}")
            blocks.append(Block("table", [rows[0], *rows[2:]]))
            continue

        if line.startswith("# "):
            blocks.append(Block("chapter", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(Block("heading2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(Block("heading3", line[4:].strip()))
        elif line.startswith("#### "):
            blocks.append(Block("heading4", line[5:].strip()))
        elif line.startswith("*Sumber:") and line.endswith("*"):
            blocks.append(Block("source", line[1:-1]))
        elif ORDERED_ITEM.match(line):
            item = ORDERED_ITEM.match(line)
            assert item is not None
            blocks.append(Block("ordered_item", item.group(2), int(item.group(1))))
        elif line.startswith("`") and line.endswith("`") and "[" in line:
            blocks.append(Block("placeholder", line.strip("`")))
        else:
            blocks.append(Block("paragraph", line))
        index += 1

    return blocks


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def set_run_font(
    run, size: float = 12, *, bold: bool | None = None, italic: bool | None = None
) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run_element = run._element
    run_properties = run_element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")
    language = run_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        run_properties.append(language)
    language.set(qn("w:val"), "id-ID")


def add_inline_runs(
    paragraph, text: str, *, size: float = 12, force_bold: bool = False, force_italic: bool = False
) -> None:
    cursor = 0
    for match in INLINE_MARKUP.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size, bold=force_bold, italic=force_italic)

        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True, italic=force_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, bold=force_bold, italic=force_italic)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, bold=force_bold, italic=True)
        cursor = match.end()

    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size, bold=force_bold, italic=force_italic)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, *, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80
) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table) -> None:
    table_properties = table._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_cm: float) -> None:
    width = Cm(width_cm)
    cell.width = width
    cell_properties = cell._tc.get_or_add_tcPr()
    tc_width = cell_properties.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        cell_properties.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width.twips)))
    tc_width.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def configure_paragraph(
    paragraph, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent: bool = True
) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Cm(1) if first_indent else Cm(0)
    fmt.widow_control = True


def add_body_paragraph(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    configure_paragraph(paragraph)
    add_inline_runs(paragraph, text)


def add_heading(document: DocumentObject, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(18 if level == 2 else 12)
    fmt.space_after = Pt(6)
    fmt.keep_with_next = True
    fmt.keep_together = True
    add_inline_runs(paragraph, text, force_bold=True)


def add_chapter_title(document: DocumentObject, text: str, ordinal: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0 if ordinal == 0 else 18)
    fmt.keep_with_next = True
    add_inline_runs(paragraph, text, force_bold=True)


def add_caption(document: DocumentObject, text: str, *, figure: bool) -> None:
    paragraph = document.add_paragraph()
    configure_paragraph(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER if figure else WD_ALIGN_PARAGRAPH.LEFT,
        first_indent=False,
    )
    paragraph.paragraph_format.space_before = Pt(3 if figure else 6)
    paragraph.paragraph_format.space_after = Pt(6 if figure else 3)
    paragraph.paragraph_format.keep_with_next = not figure
    paragraph.paragraph_format.keep_together = True
    add_inline_runs(paragraph, text)


def add_source(document: DocumentObject, text: str, *, centered: bool) -> None:
    paragraph = document.add_paragraph()
    configure_paragraph(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT,
        first_indent=False,
    )
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    add_inline_runs(paragraph, text, size=10, force_italic=True)


def add_placeholder(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    configure_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF2CC")
    properties.append(shading)
    add_inline_runs(paragraph, text, force_bold=True)


def add_ordered_item(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    configure_paragraph(paragraph, first_indent=False)
    paragraph.paragraph_format.left_indent = Cm(1)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    add_inline_runs(paragraph, text)


def add_table(document: DocumentObject, rows: list[list[str]], table_number: int) -> None:
    if not rows:
        raise ValueError(f"Tabel 4.{table_number} tidak memiliki baris")
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError(f"Jumlah kolom tidak konsisten pada Tabel 4.{table_number}")

    widths = TABLE_WIDTHS_CM[table_number]
    if len(widths) != column_count:
        raise ValueError(f"Konfigurasi lebar Tabel 4.{table_number} tidak cocok")

    font_size = 9.5 if column_count <= 4 else 9.0
    if table_number in {6, 9, 10}:
        font_size = 8.5
    if table_number in {7, 12}:
        font_size = 8.0

    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_layout(table)
    set_table_borders(table)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, "D9E2F3")

            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            fmt = paragraph.paragraph_format
            fmt.first_line_indent = Cm(0)
            fmt.left_indent = Cm(0)
            fmt.right_indent = Cm(0)
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.widow_control = True
            add_inline_runs(paragraph, value, size=font_size, force_bold=row_index == 0)

    following = document.add_paragraph()
    following.paragraph_format.space_after = Pt(0)
    following.paragraph_format.line_spacing = Pt(1)


def add_picture(document: DocumentObject, image_path: Path) -> None:
    if not image_path.is_file() or image_path.stat().st_size == 0:
        raise FileNotFoundError(f"Gambar BAB IV tidak valid: {image_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(14.0))


def set_page_size(section, landscape: bool) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    section.left_margin = Cm(4.0)
    section.top_margin = Cm(4.0)
    section.right_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run()
    set_run_font(run, 12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_styles(document: DocumentObject) -> None:
    for style_name in ("Normal", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style_element = style.element
        style_properties = style_element.get_or_add_rPr()
        fonts = style_properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style_properties.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), "Times New Roman")


def table_number_after(blocks: list[Block], start: int) -> int | None:
    for block in blocks[start + 1 :]:
        if block.kind == "table":
            return int(block.meta) if block.meta is not None else None
        if block.kind == "table_caption":
            return None
    return None


def attach_table_numbers(blocks: list[Block]) -> list[Block]:
    result: list[Block] = []
    pending_number: int | None = None
    for block in blocks:
        if block.kind == "table_caption":
            pending_number = int(block.meta)
            result.append(block)
        elif block.kind == "table":
            if pending_number is None:
                raise ValueError("Tabel ditemukan tanpa caption bernomor")
            result.append(Block("table", block.value, pending_number))
            pending_number = None
        else:
            result.append(block)
    return result


def iter_blocks_with_orientation(blocks: list[Block]) -> Iterable[tuple[Block, bool, bool]]:
    in_landscape = False
    pending_wide = False
    for block in blocks:
        switch_before = False
        switch_after = False
        if block.kind == "table_caption":
            pending_wide = int(block.meta) in WIDE_TABLES
            if pending_wide and not in_landscape:
                switch_before = True
                in_landscape = True
        elif block.kind == "source" and pending_wide:
            switch_after = True
            in_landscape = False
            pending_wide = False

        yield block, switch_before, switch_after


def build_document(source: Path, output: Path, *, max_blocks: int | None = None) -> dict[str, int]:
    blocks = attach_table_numbers(parse_markdown(source))
    if max_blocks is not None:
        if max_blocks < 1:
            raise ValueError("max_blocks harus lebih besar dari nol")
        blocks = blocks[:max_blocks]
    document = Document()
    configure_styles(document)
    set_page_size(document.sections[0], landscape=False)

    properties = document.core_properties
    properties.title = THESIS_TITLE
    properties.subject = "BAB IV - Hasil dan Pembahasan"
    properties.author = "Peneliti TaniAdaptif"
    properties.comments = (
        "Draft pra-E2E; placeholder bukti GPU dan screenshot Android dipertahankan."
    )

    chapter_count = 0
    table_count = 0
    figure_count = 0
    source_for_figure = False

    for block, switch_before, switch_after in iter_blocks_with_orientation(blocks):
        if switch_before:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            set_page_size(section, landscape=True)

        if block.kind == "chapter":
            add_chapter_title(document, str(block.value), chapter_count)
            chapter_count += 1
        elif block.kind == "heading2":
            add_heading(document, str(block.value), 2)
        elif block.kind == "heading3":
            add_heading(document, str(block.value), 3)
        elif block.kind == "heading4":
            add_heading(document, str(block.value), 4)
        elif block.kind == "paragraph":
            add_body_paragraph(document, str(block.value))
        elif block.kind == "placeholder":
            add_placeholder(document, str(block.value))
        elif block.kind == "ordered_item":
            add_ordered_item(document, str(block.value))
        elif block.kind == "table_caption":
            add_caption(document, str(block.value), figure=False)
        elif block.kind == "table":
            table_number = int(block.meta)
            add_table(document, block.value, table_number)  # type: ignore[arg-type]
            table_count += 1
            source_for_figure = False
        elif block.kind == "image":
            image_path = (source.parent / str(block.value)).resolve()
            add_picture(document, image_path)
            figure_count += 1
            source_for_figure = True
        elif block.kind == "figure_caption":
            add_caption(document, str(block.value), figure=True)
        elif block.kind == "source":
            add_source(document, str(block.value), centered=source_for_figure)
            source_for_figure = False
        else:
            raise ValueError(f"Jenis blok tidak dikenal: {block.kind}")

        if switch_after:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            set_page_size(section, landscape=False)

    for index, section in enumerate(document.sections):
        if index > 0:
            section.footer.is_linked_to_previous = True
        else:
            add_page_number(section)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.docx")
    document.save(temporary)
    temporary.replace(output)

    return {
        "blocks": len(blocks),
        "tables": table_count,
        "figures": figure_count,
        "sections": len(document.sections),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Bangun DOCX BAB IV dari draft Markdown tervalidasi."
    )
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--max-blocks",
        type=int,
        default=None,
        help="argumen diagnostik untuk render parsial",
    )
    args = parser.parse_args()

    source = args.source.resolve()
    output = args.output.resolve()
    if not source.is_file():
        raise SystemExit(f"Draft BAB IV tidak ditemukan: {source}")

    result = build_document(source, output, max_blocks=args.max_blocks)
    if args.max_blocks is None and (result["tables"] != 13 or result["figures"] != 6):
        raise SystemExit(f"Jumlah artefak DOCX tidak sesuai: {result}")
    print(f"DOCX: {output}")
    print(f"Ukuran: {output.stat().st_size:,} byte")
    print(f"Ringkasan: {result}")


if __name__ == "__main__":
    main()
