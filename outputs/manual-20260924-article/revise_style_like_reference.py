from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"D:\TaniAdaptif\outputs\manual-20260924-article\Artikel_Ilmiah_TaniAdaptif_Infotek_Revisi_Pendahuluan_Landasan.docx")
OUT = Path(r"D:\TaniAdaptif\outputs\manual-20260924-article\Artikel_Ilmiah_TaniAdaptif_Infotek_Revisi_Gaya_Artikel.docx")


def set_text(paragraph, text, size=12):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    run.font.name = "Arial Narrow"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial Narrow")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial Narrow")
    run.font.size = Pt(size)


doc = Document(SRC)
paragraphs = list(doc.paragraphs)

intro_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Pendahuluan")
intro_end = next(i for i in range(intro_heading + 1, len(paragraphs)) if paragraphs[i].text.strip() == "Tinjauan Pustaka")
intro_body = [p for p in paragraphs[intro_heading + 1 : intro_end] if p.text.strip()]
intro_text = [
    "Sektor pertanian memiliki hubungan yang erat dengan kondisi cuaca dan iklim. Perubahan curah hujan, ketersediaan air, dan keadaan lingkungan dapat memengaruhi waktu tanam serta pilihan tanaman. Karena itu, penyesuaian pola tanam diperlukan agar kegiatan budidaya dapat mempertimbangkan kondisi lingkungan yang berubah. Rotasi tanam menjadi salah satu cara untuk menyusun tanaman secara bergantian pada lahan yang sama dan mendukung adaptasi terhadap perubahan iklim [1], [17].",
    "Perkembangan teknologi informasi memungkinkan informasi cuaca diperoleh melalui layanan digital. BMKG menyediakan data cuaca berdasarkan wilayah yang dapat digunakan oleh aplikasi. Data terbuka BMKG dapat diakses dan diintegrasikan melalui mekanisme pertukaran data [18]. Dengan demikian, informasi cuaca tidak hanya ditampilkan kepada pengguna, tetapi juga dapat diproses sebagai salah satu masukan dalam penyusunan rekomendasi tanaman.",
    "Penelitian sebelumnya telah memanfaatkan aplikasi Android untuk menyajikan informasi cuaca dalam bidang pertanian [3]. Penelitian lain menggunakan sistem pendukung keputusan untuk menentukan pilihan tanaman berdasarkan beberapa kriteria [4], [12]. Pendekatan tersebut menunjukkan bahwa teknologi informasi dapat membantu pengolahan informasi pertanian, tetapi belum menggabungkan data cuaca BMKG, kondisi lahan, tanaman sebelumnya, rule-based engine, dan visualisasi generatif dalam satu alur aplikasi.",
    "Selain informasi dalam bentuk teks, hasil rekomendasi dapat disajikan melalui simulasi visual. Stable Diffusion merupakan model generatif yang dapat menghasilkan gambar berdasarkan prompt dan dapat disesuaikan melalui proses fine-tuning [5]. LoRA digunakan sebagai teknik penyesuaian model dengan parameter yang lebih ringan [6]. Pada penelitian ini, Stable Diffusion 1.5 dan LoRA digunakan setelah rekomendasi diperoleh. Keduanya tidak digunakan untuk menentukan tanaman, melainkan untuk menghasilkan visualisasi komoditas yang telah dipilih oleh rule-based engine.",
    "Berdasarkan permasalahan tersebut, penelitian ini mengembangkan TaniAdaptif sebagai aplikasi mobile Android untuk simulasi rotasi tanam adaptif iklim di wilayah Aikmel. Sistem mengintegrasikan data cuaca BMKG, jenis lahan, ketersediaan air, tanaman sebelumnya, dan rule-based engine berbasis decision table. Hasil rekomendasi kemudian dapat divisualisasikan menggunakan Stable Diffusion 1.5 dengan fine-tuning LoRA pada empat komoditas, yaitu padi, jagung, cabai, dan kacang tanah. Penelitian berfokus pada perancangan, implementasi, dan pengujian teknis prototipe, bukan pada validasi agronomi lapangan atau prediksi hasil panen.",
]
for paragraph, text in zip(intro_body, intro_text):
    set_text(paragraph, text)
if len(intro_body) < len(intro_text):
    new_paragraph = doc.add_paragraph(style="Body")
    set_text(new_paragraph, intro_text[-1])
    intro_body[-1]._p.addnext(new_paragraph._p)

landasan_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Landasan Konsep")
method_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Metode Penelitian")
concept_paragraphs = [
    p for p in paragraphs[landasan_heading + 1 : method_heading]
    if p.style.name == "Body"
]
concept_text = [
    "Rotasi tanam merupakan pola penanaman beberapa jenis tanaman secara bergantian pada lahan yang sama dalam periode tanam yang berbeda. Dalam penelitian ini, rotasi tanam digunakan untuk menyusun rekomendasi tanaman berikutnya dengan mempertimbangkan cuaca, jenis lahan, ketersediaan air, dan tanaman sebelumnya. Konsep ini digunakan pada empat komoditas, yaitu padi, jagung, cabai, dan kacang tanah, serta tidak digunakan untuk memprediksi hasil panen [1], [17].",
    "BMKG merupakan sumber data cuaca yang digunakan dalam sistem. Data cuaca berdasarkan wilayah diambil dan dinormalisasi sebelum diproses bersama kondisi lahan. Informasi tersebut menjadi masukan untuk membantu sistem menyusun rekomendasi. Pemanfaatan data BMKG dalam aplikasi sejalan dengan penggunaan data terbuka untuk mendukung layanan informasi berbasis perangkat lunak [3], [18].",
    "Rule-based engine adalah pendekatan yang menghasilkan keluaran berdasarkan aturan yang telah ditentukan. Pada TaniAdaptif, aturan disusun dalam bentuk decision table agar kombinasi kondisi masukan dapat dipetakan ke rekomendasi tanaman. Setiap keluaran dapat ditelusuri melalui kunci aturan, alasan, dan rule ID. Pendekatan ini digunakan karena proses keputusan perlu tetap terstruktur dan dapat diaudit [4], [12].",
    "Stable Diffusion merupakan model generatif text-to-image yang menghasilkan gambar berdasarkan prompt. Fine-tuning digunakan untuk menyesuaikan model yang telah tersedia dengan karakteristik dataset tertentu, sedangkan LoRA digunakan sebagai teknik penyesuaian dengan jumlah parameter yang lebih ringan [5], [6], [21], [22]. Dalam penelitian ini, Stable Diffusion dan LoRA hanya berfungsi sebagai pembangkit simulasi visual setelah rule-based engine menentukan komoditas.",
    "Black Box Testing merupakan pengujian yang berfokus pada masukan, keluaran, dan fungsi sistem tanpa memeriksa struktur internal program. Metode ini digunakan untuk memeriksa fungsi utama TaniAdaptif, seperti pengolahan cuaca, validasi input, pemrosesan aturan, layanan inference, penyimpanan, dan pengiriman gambar. Pengujian ini menjadi bukti fungsional dan teknis prototipe, bukan UAT atau validasi agronomi [25].",
]
for paragraph, text in zip(concept_paragraphs, concept_text):
    set_text(paragraph, text)

doc.core_properties.comments = "Revisi gaya terbatas pada Pendahuluan dan Landasan Konsep; fakta bersumber dari skripsi TaniAdaptif."
doc.save(OUT)
print(OUT)
