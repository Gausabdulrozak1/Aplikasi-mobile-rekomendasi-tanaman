from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"D:\TaniAdaptif\outputs\manual-20260924-article\Artikel_Ilmiah_TaniAdaptif_Infotek_Revisi_Gaya_Artikel.docx")
OUT = Path(r"D:\TaniAdaptif\outputs\manual-20260924-article\Artikel_Ilmiah_TaniAdaptif_Infotek_FINAL.docx")


def replace_text(paragraph, text, size=12):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    run.font.name = "Arial Narrow"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial Narrow")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial Narrow")
    run.font.size = Pt(size)


doc = Document(SRC)
paragraphs = list(doc.paragraphs)

# Section labels follow the reference article's presentation style.
for p in paragraphs:
    if p.text.strip() == "Pendahuluan":
        replace_text(p, "1. Pendahuluan")
    elif p.text.strip() == "Tinjauan Pustaka":
        replace_text(p, "2. Tinjauan Pustaka")
    elif p.text.strip() == "Penelitian Terkait":
        replace_text(p, "2.1 Penelitian Terkait")
    elif p.text.strip() == "Landasan Konsep":
        replace_text(p, "2.2 Landasan Teori")

# The paragraphs below use the same short definition -> application pattern as
# the reference article, while retaining TaniAdaptif's own facts and sources.
intro = [
    "Sektor pertanian merupakan salah satu bidang yang sangat dipengaruhi oleh kondisi cuaca dan iklim. Perubahan curah hujan, ketersediaan air, dan keadaan lingkungan dapat memengaruhi waktu tanam serta pilihan tanaman. Oleh karena itu, penyesuaian pola tanam perlu dilakukan agar kegiatan budidaya dapat mempertimbangkan perubahan kondisi lingkungan. Salah satu pendekatan yang dapat digunakan adalah rotasi tanam, yaitu penanaman beberapa komoditas secara bergantian pada lahan yang sama [1], [17].",
    "Perkembangan teknologi informasi memberikan peluang untuk mengolah informasi cuaca melalui aplikasi. Badan Meteorologi, Klimatologi, dan Geofisika (BMKG) menyediakan data cuaca berdasarkan wilayah yang dapat diakses melalui layanan digital. Data terbuka BMKG dapat dimanfaatkan dan diintegrasikan ke dalam sistem informasi [18]. Dengan adanya integrasi tersebut, data cuaca dapat digunakan sebagai salah satu masukan dalam penyusunan rekomendasi tanaman, bukan hanya sebagai informasi yang ditampilkan kepada pengguna.",
    "Beberapa penelitian telah memanfaatkan aplikasi Android untuk menyajikan informasi cuaca pada bidang pertanian [3]. Penelitian lain menerapkan sistem pendukung keputusan untuk membantu pemilihan tanaman berdasarkan beberapa kriteria [4], [12]. Penelitian-penelitian tersebut menunjukkan bahwa teknologi informasi dapat membantu pengelolaan informasi pertanian. Namun, masih diperlukan sistem yang menggabungkan data cuaca, kondisi lahan, tanaman sebelumnya, rekomendasi berbasis aturan, dan simulasi visual dalam satu aplikasi.",
    "Stable Diffusion merupakan model generatif yang dapat menghasilkan gambar berdasarkan deskripsi teks atau prompt. Model tersebut dapat disesuaikan melalui proses fine-tuning untuk menghasilkan visual dengan karakteristik tertentu [5]. LoRA digunakan sebagai teknik penyesuaian model dengan jumlah parameter yang lebih ringan dibandingkan pelatihan ulang seluruh model [6]. Dalam penelitian ini, Stable Diffusion 1.5 dengan LoRA digunakan sebagai komponen visualisasi setelah rekomendasi tanaman diperoleh. Model generatif tidak digunakan untuk menentukan keputusan rekomendasi.",
    "Berdasarkan permasalahan tersebut, penelitian ini mengembangkan aplikasi TaniAdaptif untuk simulasi rotasi tanam adaptif iklim di wilayah Aikmel. Aplikasi ini mengintegrasikan data cuaca BMKG, jenis lahan, ketersediaan air, tanaman sebelumnya, dan rule-based engine berbasis decision table. Rekomendasi yang dihasilkan dapat divisualisasikan menggunakan Stable Diffusion 1.5 dengan fine-tuning LoRA pada komoditas padi, jagung, cabai, dan kacang tanah. Penelitian berfokus pada pengembangan serta pengujian teknis prototipe dan tidak dimaksudkan sebagai validasi agronomi atau prediksi hasil panen.",
]

paragraphs = list(doc.paragraphs)
start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "1. Pendahuluan")
end = next(i for i in range(start + 1, len(paragraphs)) if paragraphs[i].text.strip() == "2. Tinjauan Pustaka")
body = [p for p in paragraphs[start + 1 : end] if p.text.strip()]
for p, text in zip(body, intro):
    replace_text(p, text)

theory = [
    "Rotasi tanam merupakan pola penanaman beberapa jenis tanaman secara bergantian pada lahan yang sama dalam periode tanam yang berbeda. Dalam penelitian ini, rotasi tanam digunakan untuk menyusun rekomendasi tanaman berikutnya berdasarkan kondisi cuaca, jenis lahan, ketersediaan air, dan tanaman sebelumnya. Komoditas yang digunakan dibatasi pada padi, jagung, cabai, dan kacang tanah. Rotasi tanam pada penelitian ini tidak digunakan untuk memprediksi produktivitas tanaman [1], [17].",
    "Data cuaca BMKG digunakan sebagai sumber informasi lingkungan dalam sistem. Data berdasarkan wilayah diambil, dinormalisasi, dan diproses bersama masukan kondisi lahan. Informasi tersebut kemudian diteruskan ke proses rekomendasi sehingga sistem dapat memberikan keluaran berdasarkan kondisi yang tersedia. Pemanfaatan data BMKG dalam sistem informasi mendukung pengolahan data cuaca sebagai masukan aplikasi pertanian [3], [18].",
    "Rule-based engine merupakan pendekatan yang menggunakan aturan tertentu untuk menghasilkan suatu keluaran. Pada TaniAdaptif, aturan tersebut disusun dalam bentuk decision table yang memetakan kombinasi kondisi masukan ke rekomendasi tanaman. Setiap rekomendasi dapat ditelusuri melalui kunci aturan, alasan, dan rule ID. Pendekatan ini digunakan agar proses keputusan berjalan secara terstruktur dan dapat diaudit [4], [12].",
    "Stable Diffusion adalah model generatif text-to-image yang menghasilkan gambar berdasarkan prompt. Fine-tuning digunakan untuk menyesuaikan model yang telah tersedia terhadap dataset tertentu, sedangkan LoRA digunakan untuk melakukan penyesuaian dengan parameter yang lebih ringan [5], [6], [21], [22]. Pada TaniAdaptif, Stable Diffusion dan LoRA hanya digunakan untuk menghasilkan simulasi visual dari komoditas yang telah direkomendasikan oleh rule-based engine.",
    "Black Box Testing merupakan metode pengujian yang berfokus pada masukan, keluaran, dan fungsi sistem tanpa memeriksa struktur internal program. Pengujian ini digunakan untuk memeriksa proses BMKG, validasi input, rule-based engine, layanan inference, penyimpanan, dan pengiriman gambar. Hasil pengujian menjadi bukti fungsional dan teknis prototipe, bukan bukti UAT atau validasi agronomi [25].",
]

paragraphs = list(doc.paragraphs)
landasan = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "2.2 Landasan Teori")
method = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Metode Penelitian")
concept_headings = [p for p in paragraphs[landasan + 1 : method] if p.style.name == "Heading 3"]
concept_body = [p for p in paragraphs[landasan + 1 : method] if p.style.name == "Body"]
for p, text in zip(concept_body, theory):
    replace_text(p, text)

for index, heading in enumerate(concept_headings, start=1):
    names = [
        "1. Rotasi Tanam dan Adaptasi Iklim",
        "2. Data Cuaca BMKG dan Sistem Rekomendasi",
        "3. Rule-Based Engine dan Decision Table",
        "4. Stable Diffusion dan LoRA",
        "5. Pengujian Sistem",
    ]
    replace_text(heading, names[index - 1], size=11)

doc.core_properties.comments = "Final revisi gaya artikel; hanya Pendahuluan dan Landasan Teori yang disesuaikan dengan pola artikel rujukan."
doc.save(OUT)
print(OUT)
